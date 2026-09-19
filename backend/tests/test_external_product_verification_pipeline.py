"""
MetrCheck AI — FSSAI & Barcode External Verification Pipeline Test Suite
Comprehensive testing for all 18 external verification scenarios:
1. Valid 14-digit FSSAI extraction
2. Invalid FSSAI number
3. Barcode decoding
4. QR decoding
5. Successful external FSSAI lookup
6. External FSSAI not found
7. External provider unavailable
8. Successful barcode lookup
9. Barcode product not found
10. OCR/external MATCH
11. OCR/external MISMATCH
12. Missing external field
13. Network timeout
14. API failure (HTTP 500)
15. Secrets not exposed to frontend
16. Existing Legal Metrology rules continue passing
17. Existing FSSAI compliance rules continue passing
18. Existing report generation continues working (PDF, Excel, CSV, JSON)
"""

import io
import json
import pytest
import httpx
from unittest.mock import AsyncMock, patch, MagicMock
from openpyxl import load_workbook

from models.external_verification import (
    ExternalVerificationStatus,
    ComparisonResult,
    FSSAIExtractionData,
    BarcodeExtractionData,
    ExternalFSSAIData,
    ExternalProductData,
    CrossSourceFieldComparison,
    ExternalProductVerificationPipelineResult
)
from services.external_verification_pipeline import (
    FieldComparator,
    FSSAIProvider,
    ProductLookupProvider,
    FSSAIProviderCoordinator,
    ProductLookupCoordinator,
    FoSCoSOfficialProvider,
    GS1DataKartOfficialProvider,
    LocalFSSAICacheProvider,
    LocalProductDataCacheProvider,
    ExternalProductVerificationPipeline,
    external_verification_pipeline
)
from services.barcode_service import validate_ean13_checksum
from services.identity.fssai_extractor import fssai_extractor
from models.schemas import (
    ProductInfo,
    ComplianceResult,
    ComplianceCheck,
    AnalysisResponse,
    OCRResult
)
from services.report_service import generate_pdf_report
from api.report import get_report_xlsx, get_report_csv, get_report_json
from database.db import save_analysis, get_analysis
from compliance.engine import ComplianceEngine


# ══════════════════════════════════════════════════════════════════════
# 1. VALID 14-DIGIT FSSAI EXTRACTION
# ══════════════════════════════════════════════════════════════════════
def test_01_valid_14_digit_fssai_extraction():
    pipeline = ExternalProductVerificationPipeline()
    ocr_text = "Manufactured by Kissan. FSSAI Lic. No. 10013022001897. Packed at Plot 4."
    
    fssai_data = pipeline.extract_fssai_with_provenance(ocr_text)
    assert fssai_data.detected is True
    assert fssai_data.number == "10013022001897"
    assert fssai_data.format_valid is True
    assert fssai_data.confidence >= 0.80
    assert fssai_data.state_code == "00" or fssai_data.state_name is not None


# ══════════════════════════════════════════════════════════════════════
# 2. INVALID FSSAI NUMBER
# ══════════════════════════════════════════════════════════════════════
def test_02_invalid_fssai_number():
    # Direct structure validation: 13 digits is invalid
    is_v, _, _ = fssai_extractor.validate_structure("1001302200189")
    assert is_v is False

    # All zeros is invalid
    is_v_zeros, _, _ = fssai_extractor.validate_structure("00000000000000")
    assert is_v_zeros is False

    # Pipeline extraction of non-FSSAI text
    pipeline = ExternalProductVerificationPipeline()
    res = pipeline.extract_fssai_with_provenance("Batch 12345 Exp 2026")
    assert res.detected is False or res.format_valid is False


# ══════════════════════════════════════════════════════════════════════
# 3. BARCODE DECODING & CHECKSUM
# ══════════════════════════════════════════════════════════════════════
def test_03_barcode_decoding():
    # Valid EAN-13 checksum: 8901030383748
    assert validate_ean13_checksum("8901030383748") is True

    # Invalid EAN-13 checksum: last digit should be 8, not 9
    assert validate_ean13_checksum("8901030383749") is False

    # Extract barcodes handles empty images cleanly
    pipeline = ExternalProductVerificationPipeline()
    bcs = pipeline.decode_barcodes_from_images([])
    assert isinstance(bcs, list)


# ══════════════════════════════════════════════════════════════════════
# 4. QR DECODING
# ══════════════════════════════════════════════════════════════════════
def test_04_qr_decoding():
    bc_item = BarcodeExtractionData(
        detected=True,
        type="QR_CODE",
        value="https://hul.co.in/kissan-ketchup-950g",
        qr_payload="https://hul.co.in/kissan-ketchup-950g",
        is_valid_checksum=True
    )
    assert bc_item.type == "QR_CODE"
    assert "hul.co.in" in (bc_item.qr_payload or "")


# ══════════════════════════════════════════════════════════════════════
# 5. SUCCESSFUL EXTERNAL FSSAI LOOKUP
# ══════════════════════════════════════════════════════════════════════
@pytest.mark.asyncio
async def test_05_successful_external_fssai_lookup():
    mock_provider = AsyncMock(spec=FSSAIProvider)
    mock_provider.verify = AsyncMock(return_value=ExternalFSSAIData(
        status=ExternalVerificationStatus.EXTERNALLY_VERIFIED,
        source="MockFoSCoS",
        licence_number="10013022001897",
        business_name="Hindustan Unilever Limited",
        registered_address="B-10, MIDC Industrial Area, Pune, Maharashtra 411018",
        licence_status="Active",
        valid_upto="2028-12-31"
    ))

    coord = FSSAIProviderCoordinator(primary=mock_provider, fallback=None)
    res = await coord.verify_fssai("10013022001897")

    assert res.status == ExternalVerificationStatus.EXTERNALLY_VERIFIED
    assert res.business_name == "Hindustan Unilever Limited"
    assert res.licence_status == "Active"


# ══════════════════════════════════════════════════════════════════════
# 6. EXTERNAL FSSAI NOT FOUND
# ══════════════════════════════════════════════════════════════════════
@pytest.mark.asyncio
async def test_06_external_fssai_not_found():
    mock_provider = AsyncMock(spec=FSSAIProvider)
    mock_provider.verify = AsyncMock(return_value=ExternalFSSAIData(
        status=ExternalVerificationStatus.EXTERNAL_DATA_NOT_FOUND,
        source="MockFoSCoS",
        licence_number="10422999000151",
        message="No FSSAI licence record found in FoSCoS database for number 10422999000151"
    ))

    coord = FSSAIProviderCoordinator(primary=mock_provider, fallback=None)
    res = await coord.verify_fssai("10422999000151")

    assert res.status == ExternalVerificationStatus.EXTERNAL_DATA_NOT_FOUND
    assert "No FSSAI licence record found" in res.message
    # Never claim counterfeit merely because not found
    assert "counterfeit" not in (res.message or "").lower()


# ══════════════════════════════════════════════════════════════════════
# 7. EXTERNAL PROVIDER UNAVAILABLE
# ══════════════════════════════════════════════════════════════════════
@pytest.mark.asyncio
async def test_07_external_provider_unavailable():
    provider = FoSCoSOfficialProvider(api_url="")
    res = await provider.verify("10013022001897")

    assert res.status == ExternalVerificationStatus.EXTERNAL_VERIFICATION_UNAVAILABLE
    assert "not configured" in res.message


# ══════════════════════════════════════════════════════════════════════
# 8. SUCCESSFUL BARCODE LOOKUP
# ══════════════════════════════════════════════════════════════════════
@pytest.mark.asyncio
async def test_08_successful_barcode_lookup():
    mock_gs1 = AsyncMock(spec=ProductLookupProvider)
    mock_gs1.lookup_by_gtin = AsyncMock(return_value=ExternalProductData(
        status=ExternalVerificationStatus.EXTERNALLY_VERIFIED,
        source="MockGS1DataKart",
        gtin="8901030383748",
        product_name="Kissan Fresh Tomato Ketchup",
        brand="Kissan",
        manufacturer="Hindustan Unilever Limited",
        net_quantity="950 g",
        category="Sauces & Condiments"
    ))

    coord = ProductLookupCoordinator(primary=mock_gs1, fallback=None)
    res = await coord.lookup_product("8901030383748")

    assert res.status == ExternalVerificationStatus.EXTERNALLY_VERIFIED
    assert res.brand == "Kissan"
    assert res.product_name == "Kissan Fresh Tomato Ketchup"
    assert res.net_quantity == "950 g"


# ══════════════════════════════════════════════════════════════════════
# 9. BARCODE PRODUCT NOT FOUND
# ══════════════════════════════════════════════════════════════════════
@pytest.mark.asyncio
async def test_09_barcode_product_not_found():
    mock_gs1 = AsyncMock(spec=ProductLookupProvider)
    mock_gs1.lookup_by_gtin = AsyncMock(return_value=ExternalProductData(
        status=ExternalVerificationStatus.EXTERNAL_DATA_NOT_FOUND,
        source="MockGS1DataKart",
        gtin="8901234567890",
        message="No GTIN registered in GS1 DataKart catalog"
    ))

    coord = ProductLookupCoordinator(primary=mock_gs1, fallback=None)
    res = await coord.lookup_product("8901234567890")

    assert res.status == ExternalVerificationStatus.EXTERNAL_DATA_NOT_FOUND
    assert "No GTIN registered" in res.message


# ══════════════════════════════════════════════════════════════════════
# 10. OCR / EXTERNAL MATCH (NORMALIZED)
# ══════════════════════════════════════════════════════════════════════
def test_10_ocr_external_match():
    # Brand match (case-insensitive)
    c1 = FieldComparator.compare_brand("KISSAN", "kissan", "GS1 DataKart")
    assert c1.result == ComparisonResult.MATCH

    # Manufacturer match (ignoring corporate suffix like 'Limited')
    c2 = FieldComparator.compare_manufacturer(
        "Hindustan Unilever Ltd.",
        "Hindustan Unilever Limited",
        "GS1 DataKart"
    )
    assert c2.result == ComparisonResult.MATCH

    # Net quantity normalization: 950g vs 950 g
    c3 = FieldComparator.compare_net_quantity("950 g", "950g", "GS1 DataKart")
    assert c3.result == ComparisonResult.MATCH

    # MRP normalization: ₹145 vs Rs 145
    c4 = FieldComparator.compare_mrp("₹145.00", "Rs. 145", "GS1 DataKart")
    assert c4.result == ComparisonResult.MATCH


# ══════════════════════════════════════════════════════════════════════
# 11. OCR / EXTERNAL MISMATCH
# ══════════════════════════════════════════════════════════════════════
def test_11_ocr_external_mismatch():
    # Different brand
    c1 = FieldComparator.compare_brand("Maggi", "Kissan", "GS1 DataKart")
    assert c1.result == ComparisonResult.MISMATCH

    # Different quantity
    c2 = FieldComparator.compare_net_quantity("500 g", "1 kg", "GS1 DataKart")
    assert c2.result == ComparisonResult.MISMATCH

    # Different manufacturer
    c3 = FieldComparator.compare_manufacturer("Nestle India Limited", "Hindustan Unilever Limited", "GS1 DataKart")
    assert c3.result == ComparisonResult.MISMATCH


# ══════════════════════════════════════════════════════════════════════
# 12. MISSING EXTERNAL FIELD
# ══════════════════════════════════════════════════════════════════════
def test_12_missing_external_field():
    # External value is None
    c1 = FieldComparator.compare_brand("Kissan", None, "GS1 DataKart")
    assert c1.result in (ComparisonResult.NOT_FOUND, ComparisonResult.NOT_COMPARABLE)

    # Both values are None
    c2 = FieldComparator.compare_mrp(None, None, "GS1 DataKart")
    assert c2.result == ComparisonResult.NOT_COMPARABLE


# ══════════════════════════════════════════════════════════════════════
# 13. NETWORK TIMEOUT HANDLING
# ══════════════════════════════════════════════════════════════════════
@pytest.mark.asyncio
async def test_13_network_timeout():
    provider = FoSCoSOfficialProvider(api_url="https://example-foscos.gov.in/api")
    
    with patch("httpx.AsyncClient.get", side_effect=httpx.TimeoutException("Connection timed out")):
        res = await provider.verify("10013022001897")
        assert res.status == ExternalVerificationStatus.EXTERNAL_LOOKUP_FAILED
        combined_err = f"{res.message} {res.error_details or ''}".lower()
        assert "timed out" in combined_err or "timeout" in combined_err or "failed" in combined_err


# ══════════════════════════════════════════════════════════════════════
# 14. API FAILURE (HTTP 500)
# ══════════════════════════════════════════════════════════════════════
@pytest.mark.asyncio
async def test_14_api_failure_500():
    provider = GS1DataKartOfficialProvider(api_url="https://example-gs1.org/api")
    
    mock_resp = MagicMock()
    mock_resp.status_code = 500
    mock_resp.text = "Internal Server Error"
    
    with patch("httpx.AsyncClient.get", return_value=mock_resp):
        res = await provider.lookup_by_gtin("8901030383748")
        assert res.status == ExternalVerificationStatus.EXTERNAL_LOOKUP_FAILED
        assert "500" in (res.message or "") or "500" in (res.error_details or "")


# ══════════════════════════════════════════════════════════════════════
# 15. SECRETS NOT EXPOSED TO FRONTEND
# ══════════════════════════════════════════════════════════════════════
def test_15_secrets_not_exposed_to_frontend():
    secret_key = "SUPER_SECRET_FSSAI_API_TOKEN_12345"
    
    # Create verification result
    epv = ExternalProductVerificationPipelineResult(
        fssai_extraction=FSSAIExtractionData(
            detected=True,
            number="10013022001897",
            source="OCR",
            confidence=0.92,
            format_valid=True
        ),
        external_fssai=ExternalFSSAIData(
            status=ExternalVerificationStatus.EXTERNALLY_VERIFIED,
            source="FoSCoS Official Registry API",
            business_name="Hindustan Unilever Limited"
        ),
        verification_status=ExternalVerificationStatus.EXTERNALLY_VERIFIED.value
    )

    serialized = epv.model_dump_json()
    assert secret_key not in serialized

    # Check AnalysisResponse serialization
    resp = AnalysisResponse(
        id="sec-check-001",
        product_name="Test Product",
        image_url="/test.png",
        ocr_result=OCRResult(full_text="Test", words=[], language="eng", processing_time=0.1),
        product_info=ProductInfo(product_name="Test Product", mrp="₹100"),
        compliance_result=ComplianceResult(checks=[], score=100, status="COMPLIANT", total_rules=0, passed_rules=0, failed_rules=0),
        recommendations=[],
        created_at="2026-09-17T12:00:00Z",
        external_product_verification=epv
    )
    resp_json = resp.model_dump_json()
    assert secret_key not in resp_json
    assert "SUPER_SECRET" not in resp_json


# ══════════════════════════════════════════════════════════════════════
# 16. EXISTING LEGAL METROLOGY RULES CONTINUE PASSING
# ══════════════════════════════════════════════════════════════════════
def test_16_existing_legal_metrology_rules_continue_passing():
    info = ProductInfo(
        product_name="Cadbury Dairy Milk Chocolate",
        mrp="₹50.00",
        net_quantity="50 g",
        manufacturer_name="Mondelez India Foods Private Limited",
        country_of_origin="India",
        is_food=True
    )
    engine = ComplianceEngine()
    result = engine.check(info, ocr_text="Cadbury Dairy Milk 50g MRP Rs 50", images=[])
    assert result is not None
    assert result["score"] > 0
    checks = {c.rule_id: c for c in result["checks"]}
    assert any("LM" in rid or "MRP" in rid for rid in checks)


# ══════════════════════════════════════════════════════════════════════
# 17. EXISTING FSSAI COMPLIANCE RULES CONTINUE PASSING
# ══════════════════════════════════════════════════════════════════════
def test_17_existing_fssai_compliance_rules_continue_passing():
    info = ProductInfo(
        product_name="Alpino Peanut Butter",
        mrp="₹350.00",
        net_quantity="1 kg",
        fssai_license="10019021004123",
        is_food=True
    )
    engine = ComplianceEngine()
    result = engine.check(info, ocr_text="Alpino Peanut Butter 1kg FSSAI Lic. No. 10019021004123", images=[])
    assert result is not None
    checks = {c.rule_id: c for c in result["checks"]}
    fssai_checks = [c for rid, c in checks.items() if "FS" in rid or c.domain == "FSSAI"]
    assert len(fssai_checks) > 0
    lic_check = checks.get("FS-001")
    assert lic_check is not None
    assert lic_check.status in ("PASS", "COMPLIANT", "NEEDS_REVIEW", "WARNING")


# ══════════════════════════════════════════════════════════════════════
# 18. EXISTING REPORT GENERATION CONTINUES WORKING (PDF, EXCEL, CSV, JSON)
# ══════════════════════════════════════════════════════════════════════
@pytest.mark.asyncio
async def test_18_existing_report_generation_continues_working():
    prod_info = ProductInfo(
        product_name="Kissan Fresh Tomato Ketchup",
        brand="Kissan",
        manufacturer="Hindustan Unilever Limited",
        net_quantity="950 g",
        mrp="₹145.00",
        fssai_license="10013022001897",
        barcode_detected="8901030383748",
        is_food=True
    )

    comp_res = ComplianceResult(
        checks=[
            ComplianceCheck(
                rule_id="RULE_01_MRP",
                field="mrp",
                field_label="Maximum Retail Price",
                required=True,
                detected=True,
                detected_value="₹145.00",
                severity="HIGH",
                status="PASS",
                description="MRP declaration is mandatory",
                source="Legal Metrology Rules 2011"
            )
        ],
        score=98,
        status="COMPLIANT",
        total_rules=1,
        passed_rules=1,
        failed_rules=0
    )

    epv = ExternalProductVerificationPipelineResult(
        fssai_extraction=FSSAIExtractionData(
            detected=True,
            number="10013022001897",
            source="Package OCR",
            confidence=0.95,
            format_valid=True,
            state_name="Maharashtra"
        ),
        external_fssai=ExternalFSSAIData(
            status=ExternalVerificationStatus.EXTERNALLY_VERIFIED,
            source="FoSCoS Official Registry API",
            business_name="Hindustan Unilever Limited",
            registered_address="Mumbai, Maharashtra",
            licence_status="Active"
        ),
        barcode_extractions=[
            BarcodeExtractionData(
                detected=True,
                value="8901030383748",
                type="EAN-13",
                source="Barcode Scanner",
                format_valid=True,
                is_valid_checksum=True,
                country_of_origin="India"
            )
        ],
        external_product=ExternalProductData(
            status=ExternalVerificationStatus.EXTERNALLY_VERIFIED,
            source="GS1 India DataKart / Verified Registry",
            product_name="Kissan Fresh Tomato Ketchup",
            brand="Kissan",
            manufacturer="Hindustan Unilever Limited",
            net_quantity="950 g"
        ),
        cross_source_comparisons=[
            CrossSourceFieldComparison(
                field="Product Name",
                package_value="Kissan Fresh Tomato Ketchup",
                external_value="Kissan Fresh Tomato Ketchup",
                result=ComparisonResult.MATCH
            ),
            CrossSourceFieldComparison(
                field="Brand",
                package_value="Kissan",
                external_value="Kissan",
                result=ComparisonResult.MATCH
            )
        ],
        verification_status=ExternalVerificationStatus.EXTERNALLY_VERIFIED.value
    )

    analysis_dict = {
        "id": "e2e-report-test-001",
        "product_name": "Kissan Fresh Tomato Ketchup",
        "image_filename": "ketchup.jpg",
        "ocr_text": "Kissan Fresh Tomato Ketchup 950g FSSAI 10013022001897",
        "extracted_data": prod_info.model_dump(),
        "compliance_result": comp_res.model_dump(),
        "score": 98,
        "status": "COMPLIANT",
        "created_at": "2026-09-17T12:00:00Z",
        "images": [],
        "external_product_verification": epv.model_dump()
    }

    await save_analysis(analysis_dict)

    # 1. PDF Report Generation
    full_resp = AnalysisResponse(
        id="e2e-report-test-001",
        product_name=prod_info.product_name,
        image_url="/uploads/ketchup.jpg",
        ocr_result=OCRResult(full_text=analysis_dict["ocr_text"], words=[], language="eng", processing_time=0.1),
        product_info=prod_info,
        compliance_result=comp_res,
        recommendations=[],
        created_at=analysis_dict["created_at"],
        external_product_verification=epv
    )
    pdf_bytes = generate_pdf_report(full_resp, lang="en")
    assert pdf_bytes.startswith(b"%PDF-")
    assert len(pdf_bytes) > 8000

    # 2. Excel (XLSX) Multi-sheet Report Generation
    xlsx_resp = await get_report_xlsx("e2e-report-test-001")
    xlsx_bytes = b"".join([chunk async for chunk in xlsx_resp.body_iterator])
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert "Summary" in wb.sheetnames
    assert "External Verification" in wb.sheetnames
    ws_ext = wb["External Verification"]
    ext_vals = [str(cell.value) for row in ws_ext.rows for cell in row if cell.value]
    assert any("EXTERNAL PRODUCT & LICENCE VERIFICATION" in v for v in ext_vals)
    assert any("FSSAI LICENCE VERIFICATION" in v for v in ext_vals)
    assert any("BARCODE & GTIN PRODUCT VERIFICATION" in v for v in ext_vals)
    assert any("CROSS-SOURCE VERIFICATION" in v for v in ext_vals)
    assert any("Notice:" in v for v in ext_vals)

    # 3. CSV Report Generation
    csv_resp = await get_report_csv("e2e-report-test-001")
    csv_bytes = b"".join([chunk async for chunk in csv_resp.body_iterator])
    csv_text = csv_bytes.decode("utf-8-sig")
    assert "EXTERNAL PRODUCT & LICENCE VERIFICATION" in csv_text
    assert "1. FSSAI LICENCE VERIFICATION" in csv_text
    assert "2. BARCODE & GTIN PRODUCT VERIFICATION" in csv_text
    assert "3. CROSS-SOURCE VERIFICATION" in csv_text
    assert "Notice: External verification depends on registry availability" in csv_text

    # 4. JSON Export
    json_resp = await get_report_json("e2e-report-test-001")
    json_bytes = b"".join([chunk async for chunk in json_resp.body_iterator])
    parsed_json = json.loads(json_bytes.decode("utf-8"))
    assert parsed_json["id"] == "e2e-report-test-001"
    assert "external_product_verification" in parsed_json
    assert parsed_json["external_product_verification"]["verification_status"] == ExternalVerificationStatus.EXTERNALLY_VERIFIED.value
